#!/usr/bin/env python3
"""Create an offline Excel report from an MT5 Account History HTML export.

Run from the repository root:
    python backend/scripts/report_mt5_html.py ReportHistory-213922177.html

The script reads only the supplied HTML file. It does not connect to MT5, load
API keys, alter the database, or upload the generated workbook.
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RAW_HEADERS = {
    "positions": ["Time", "Position", "Symbol", "Type", "Volume", "Price", "S / L", "T / P", "Time", "Price", "Commission", "Swap", "Profit"],
    "orders": ["Open Time", "Order", "Symbol", "Type", "Volume", "Price", "S / L", "T / P", "Time", "State", "Comment"],
    "deals": ["Time", "Deal", "Symbol", "Type", "Direction", "Volume", "Price", "Order", "Commission", "Fee", "Swap", "Profit", "Balance", "Comment"],
    "open_positions": ["Time", "Position", "Symbol", "Type", "Volume", "Price", "S / L", "T / P", "Market Price", "Swap", "Profit", "Comment"],
}

CANONICAL_HEADERS = {
    "positions": ["Open time", "Position", "Symbol", "Side", "Volume", "Entry price", "S / L", "T / P", "Close time", "Exit price", "Commission", "Swap", "Profit"],
    "orders": RAW_HEADERS["orders"],
    "deals": RAW_HEADERS["deals"],
    "open_positions": RAW_HEADERS["open_positions"],
}


def table_rows(soup: BeautifulSoup, section: str) -> list[dict[str, str]]:
    section_key = {"Positions": "positions", "Orders": "orders", "Deals": "deals", "Open Positions": "open_positions"}[section]
    trs = soup.find_all("tr")
    start = next((i for i, tr in enumerate(trs) if tr.get_text(" ", strip=True) == section), None)
    if start is None:
        return []
    raw_header = RAW_HEADERS[section_key]
    header = CANONICAL_HEADERS[section_key]
    result: list[dict[str, str]] = []
    for tr in trs[start + 1 :]:
        if tr.find("th"):
            break
        cells = tr.find_all("td", recursive=False)
        hidden_comment = next((cell.get_text(" ", strip=True) for cell in cells if "hidden" in cell.get("class", [])), "")
        values = [cell.get_text(" ", strip=True) for cell in cells if "hidden" not in cell.get("class", [])]
        if values == raw_header:
            continue
        if len(values) != len(header):
            continue
        parsed = dict(zip(header, values))
        if hidden_comment:
            parsed["Comment"] = hidden_comment
        result.append(parsed)
    return result


def number(value: str | None) -> float | None:
    if value is None:
        return None
    normalized = value.strip().replace("\u00a0", "").replace(" ", "").replace(",", "")
    if not normalized or normalized in {"-", "—"}:
        return None
    try:
        return float(normalized)
    except ValueError:
        return None


def parse_time(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y.%m.%d %H:%M:%S")
    except ValueError:
        return None


def prompt_id(comment: str) -> int | None:
    match = re.search(r"\[AUTOPILOT\]\s*P(\d+)\b", comment or "", re.I)
    return int(match.group(1)) if match else None


def make_trade(row: dict[str, str], open_trade: bool = False) -> dict[str, Any]:
    opened = parse_time(row.get("Time") if open_trade else row.get("Open time"))
    if open_trade:
        comment = row.get("Comment", "")
        profit = number(row.get("Profit"))
        return {
            "Prompt": prompt_id(comment), "Position ticket": row.get("Position"),
            "Symbol": row.get("Symbol"), "Side": row.get("Type"),
            "Volume": number(row.get("Volume")), "Opened at (MT5 time)": opened,
            "Open date": opened.date().isoformat() if opened else "",
            "Entry price": number(row.get("Price")), "SL": number(row.get("S / L")),
            "TP": number(row.get("T / P")), "Market price": number(row.get("Market Price")),
            "Floating P&L (not realized)": profit, "Comment": comment,
            "Lifecycle status": "Open at export time",
        }
    comment = row.get("Comment", "")
    closed = parse_time(row.get("Close time", ""))
    pnl = number(row.get("Profit"))
    return {
        "Prompt": prompt_id(comment), "Position ticket": row.get("Position"),
        "Symbol": row.get("Symbol"), "Side": row.get("Type"),
        "Volume": number(row.get("Volume")), "Opened at (MT5 time)": opened,
        "Open date": opened.date().isoformat() if opened else "",
        "Entry price": number(row.get("Entry price")), "SL": number(row.get("S / L")),
        "TP": number(row.get("T / P")), "Closed at (MT5 time)": closed,
        "Close date": closed.date().isoformat() if closed else "",
        "Exit price": number(row.get("Exit price")),
        "Commission": number(row.get("Commission")), "Swap": number(row.get("Swap")),
        "Realized P&L": pnl, "Result": "Win" if pnl and pnl > 0 else "Loss" if pnl and pnl < 0 else "Breakeven/unknown",
        "Comment": comment,
    }


def write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([row.get(h) for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        values = [len(str(c.value or "")) for c in list(col)[:300]]
        ws.column_dimensions[letter].width = min(max(max(values, default=10) + 2, 12), 42)


def metrics(trades: list[dict[str, Any]]) -> dict[str, Any]:
    pnls = [t["Realized P&L"] for t in trades if t.get("Realized P&L") is not None]
    wins = [p for p in pnls if p > 0]
    losses = [p for p in pnls if p < 0]
    gross_loss = -sum(losses)
    return {
        "Closed trades": len(pnls), "Wins": len(wins), "Losses": len(losses),
        "Breakeven/unknown": len(pnls) - len(wins) - len(losses),
        "Win rate %": round(len(wins) / len(pnls) * 100, 2) if pnls else None,
        "Net realized P&L": round(sum(pnls), 2),
        "Profit factor": round(sum(wins) / gross_loss, 3) if gross_loss else ("inf" if wins else None),
        "Avg P&L / trade": round(sum(pnls) / len(pnls), 2) if pnls else None,
    }


def aggregate(trades: list[dict[str, Any]], group_key: str) -> list[dict[str, Any]]:
    groups: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for trade in trades:
        key = trade.get(group_key)
        if key not in (None, ""):
            groups[key].append(trade)
    output = []
    for key, group in groups.items():
        row = {group_key: key, **metrics(group)}
        row["Evidence"] = "Preliminary (<10 closed trades)" if len(group) < 10 else "10+ closed trades; still assess stability"
        output.append(row)
    return sorted(output, key=lambda r: (r[group_key] if isinstance(r[group_key], str) else 0))


def prompt_day_aggregate(trades: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for t in trades:
        day, prompt = t.get("Open date"), t.get("Prompt")
        if day and prompt is not None:
            groups[(day, prompt)].append(t)
    out = []
    for (day, prompt), group in sorted(groups.items()):
        closed = [t for t in group if t.get("Realized P&L") is not None]
        row = {"Entry date": day, "Prompt": prompt, "Trades opened": len(group),
               "Still open at export": len(group) - len(closed), **metrics(closed)}
        out.append(row)
    return out


def build_report(source: Path, output: Path) -> Path:
    soup = BeautifulSoup(source.read_text(encoding="utf-16", errors="replace"), "html.parser")
    positions = table_rows(soup, "Positions")
    open_positions = table_rows(soup, "Open Positions")
    orders = table_rows(soup, "Orders")
    deals = table_rows(soup, "Deals")
    all_trades = [make_trade(row) for row in positions]
    autopilot_trades = [t for t in all_trades if t["Prompt"] is not None]
    open_trades = [make_trade(row, open_trade=True) for row in open_positions]
    autopilot_open = [t for t in open_trades if t["Prompt"] is not None]

    # Build diagnostic lookups from order and deal comments. We only use a
    # unique position-specific SL/TP clue; otherwise explicitly leave unknown.
    close_clues: dict[str, set[str]] = defaultdict(set)
    for order in orders:
        comment = order.get("Comment", "").strip().lower()
        if re.match(r"\[(sl|tp)\b", comment):
            close_clues[order.get("Order", "")].add("Stop loss" if comment.startswith("[sl") else "Take profit")
    for deal in deals:
        comment = deal.get("Comment", "").strip().lower()
        if re.match(r"\[(sl|tp)\b", comment):
            close_clues[deal.get("Order", "")].add("Stop loss" if comment.startswith("[sl") else "Take profit")
    # Position section rows don't expose the closing order ticket. Preserve
    # the clue totals separately; don't guess-link a close reason to a trade.
    exit_clue_counts = defaultdict(int)
    for clues in close_clues.values():
        for clue in clues:
            exit_clue_counts[clue] += 1

    # Daily realized P&L is grouped by close date; opened counts are distinct.
    daily_closed = aggregate(autopilot_trades, "Close date")
    daily_opened = aggregate(autopilot_trades, "Open date")
    prompt_overall = aggregate(autopilot_trades, "Prompt")
    prompt_overall.sort(key=lambda r: (-(r.get("Net realized P&L") or 0), -(r.get("Closed trades") or 0)))
    daily_prompt = prompt_day_aggregate(autopilot_trades)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary_rows = [
        {"Metric": "Source file", "Value": source.name},
        {"Metric": "Account", "Value": "MT5 HTML export (account details are in source file)"},
        {"Metric": "Date basis", "Value": "MT5 terminal/broker time shown in export; no timezone conversion"},
        {"Metric": "Closed positions in export", "Value": len(positions)},
        {"Metric": "Autopilot closed trades with prompt tag", "Value": len(autopilot_trades)},
        {"Metric": "Open positions in export", "Value": len(open_positions)},
        {"Metric": "Autopilot open trades with prompt tag", "Value": len(autopilot_open)},
        {"Metric": "Prompt IDs represented", "Value": len({t['Prompt'] for t in autopilot_trades})},
        {"Metric": "SL-tagged close orders/deals (clue count; not trade-linked)", "Value": exit_clue_counts.get("Stop loss", 0)},
        {"Metric": "TP-tagged close orders/deals (clue count; not trade-linked)", "Value": exit_clue_counts.get("Take profit", 0)},
        {"Metric": "Exit reason caveat", "Value": "Positions rows have no closing order/deal ticket; SL/TP clues are counted separately, not assigned per trade."},
        {"Metric": "Prompt proof caveat", "Value": "Small samples are preliminary; historical results do not establish future performance."},
    ]
    write_sheet(summary, ["Metric", "Value"], summary_rows)
    summary.freeze_panes = "A2"

    daily_headers = ["Close date", "Closed trades", "Wins", "Losses", "Breakeven/unknown", "Win rate %", "Net realized P&L", "Profit factor", "Avg P&L / trade", "Evidence"]
    write_sheet(wb.create_sheet("Daily Close Results"), daily_headers, daily_closed)
    write_sheet(wb.create_sheet("Trades Opened by Day"), ["Open date", "Closed trades", "Wins", "Losses", "Breakeven/unknown", "Win rate %", "Net realized P&L", "Profit factor", "Avg P&L / trade", "Evidence"], daily_opened)
    write_sheet(wb.create_sheet("Prompt by Entry Day"), ["Entry date", "Prompt", "Trades opened", "Still open at export", "Closed trades", "Wins", "Losses", "Breakeven/unknown", "Win rate %", "Net realized P&L", "Profit factor", "Avg P&L / trade"], daily_prompt)
    write_sheet(wb.create_sheet("Prompt Ranking"), ["Prompt", "Closed trades", "Wins", "Losses", "Breakeven/unknown", "Win rate %", "Net realized P&L", "Profit factor", "Avg P&L / trade", "Evidence"], prompt_overall)

    trade_headers = ["Prompt", "Position ticket", "Symbol", "Side", "Volume", "Opened at (MT5 time)", "Open date", "Entry price", "SL", "TP", "Closed at (MT5 time)", "Close date", "Exit price", "Commission", "Swap", "Realized P&L", "Result", "Comment"]
    write_sheet(wb.create_sheet("Closed Autopilot Trades"), trade_headers, autopilot_trades)
    open_headers = ["Prompt", "Position ticket", "Symbol", "Side", "Volume", "Opened at (MT5 time)", "Open date", "Entry price", "SL", "TP", "Market price", "Floating P&L (not realized)", "Comment", "Lifecycle status"]
    write_sheet(wb.create_sheet("Open Autopilot Trades"), open_headers, autopilot_open)
    write_sheet(wb.create_sheet("Unmatched Closed Positions"), trade_headers, [t for t in all_trades if t["Prompt"] is None])
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Build an offline Excel report from MT5 Account History HTML.")
    parser.add_argument("html", type=Path, help="Path to the MT5 HTML report export")
    parser.add_argument("--output", type=Path, help="Output workbook path (default: alongside HTML)")
    args = parser.parse_args()
    source = args.html.expanduser().resolve()
    if not source.is_file():
        parser.error(f"Input HTML does not exist: {source}")
    output = args.output.expanduser().resolve() if args.output else source.with_name(f"{source.stem}_prompt_performance.xlsx")
    print(f"Created: {build_report(source, output)}")


if __name__ == "__main__":
    main()
